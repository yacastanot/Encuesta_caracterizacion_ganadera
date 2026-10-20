"""Reemplaza, en el Cuadro 2 ("Municipios sin información asociada") de los
3 libros, el texto descriptivo largo de la columna "Observación" (ej.
"Menos del 80% de cobertura de encuesta sobre el RUV") por el código corto
("n.d."/"o.c.") que trae `Excluidos/Municipios_Sin_Información.xlsx` - a
pedido del usuario 2026-10-20, para que Cuadro 2 use la misma convención de
código corto que ya se aplicó al resto de los cuadros (ver
`aplicar_codigos_exclusion.py`).

Cuadro 2 es idéntico en los 3 libros (un solo pendiente "cuadro2_ganadero",
ver `main.py`): mismas 88 filas, columnas E=Observación Ciclo 1,
F=Observación Ciclo 2, desde la fila 11.

IMPORTANTE: igual que `aplicar_codigos_exclusion.py`, este parche debe
volver a correrse si el libro se reensambla (el reensamblado reconstruye
Cuadro 2 desde el pendiente, que trae el texto largo original, no el código
corto).

Uso: `python -m scripts.cuadros_ganaderia.aplicar_codigos_cuadro2`"""
from __future__ import annotations

from pathlib import Path

import openpyxl

from . import config

RUTA_EXCLUSION = config.BASE_DIR / "Excluidos" / "Municipios_Sin_Información.xlsx"
FILA_INICIO = 11
COL_C_COD_MUNICIPIO = 3
COL_E_OBS_C1 = 5
COL_F_OBS_C2 = 6


def _cargar_mapa_exclusion() -> dict[str, dict[str, str]]:
    wb = openpyxl.load_workbook(RUTA_EXCLUSION, data_only=True)
    ws = wb["Cuadro 2"]
    mapa: dict[str, dict[str, str]] = {}
    for r in range(FILA_INICIO, ws.max_row + 1):
        cod_mun = ws.cell(row=r, column=3).value
        if not cod_mun:
            continue
        cod_mun = str(cod_mun).strip()
        obs_c1 = ws.cell(row=r, column=5).value
        obs_c2 = ws.cell(row=r, column=6).value
        entrada = {}
        if obs_c1:
            entrada["C1"] = str(obs_c1).strip()
        if obs_c2:
            entrada["C2"] = str(obs_c2).strip()
        if entrada:
            mapa[cod_mun] = entrada
    return mapa


def _aplicar_en_archivo(ruta: Path, mapa_exclusion: dict[str, dict[str, str]]) -> int:
    wb = openpyxl.load_workbook(ruta)
    ws = wb["Cuadro 2"]
    total = 0
    for r in range(FILA_INICIO, ws.max_row + 1):
        cod_mun = ws.cell(row=r, column=COL_C_COD_MUNICIPIO).value
        if not cod_mun:
            continue
        cod_mun = str(cod_mun).strip()
        entrada = mapa_exclusion.get(cod_mun, {})
        for col, ciclo in [(COL_E_OBS_C1, "C1"), (COL_F_OBS_C2, "C2")]:
            codigo = entrada.get(ciclo)
            celda = ws.cell(row=r, column=col)
            if codigo:
                if celda.value != codigo:
                    celda.value = codigo
                    total += 1
            elif celda.value is not None:
                # Ese ciclo no tiene código en el archivo de referencia (el
                # municipio no está excluido en ese ciclo) - se deja en blanco,
                # igual que ya hacía `comun.generar_cuadro2` (None, no texto).
                celda.value = None
                total += 1
    wb.save(ruta)
    return total


def main() -> None:
    mapa_exclusion = _cargar_mapa_exclusion()
    print(f"Municipios con código de exclusión: {len(mapa_exclusion)}")

    archivos = {
        "inventario": [
            config.OUTPUT_DIR / "Cuadros caracterización inventario Ciclos 1 y 2_2025_generado.xlsx",
            config.TEMPLATE_INVENTARIO,
        ],
        "ganadero": [
            config.OUTPUT_DIR / "Cuadros caracterización del ganadero Ciclos 1 y 2_2025_generado.xlsx",
            config.TEMPLATE_GANADERO,
        ],
        "predio_ganadero": [
            config.OUTPUT_DIR / "Cuadros Caracterización predio-ganadero Ciclos 1 y 2_2025_generado.xlsx",
            config.TEMPLATE_PREDIO_GANADERO,
        ],
    }
    for libro, rutas in archivos.items():
        print(f"\n{libro}:")
        for ruta in rutas:
            if not Path(ruta).exists():
                print(f"  (no existe, se omite: {ruta})")
                continue
            n = _aplicar_en_archivo(Path(ruta), mapa_exclusion)
            print(f"  {Path(ruta).name}: {n} celdas actualizadas")


if __name__ == "__main__":
    main()
