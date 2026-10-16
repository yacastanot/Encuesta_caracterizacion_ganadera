"""Reemplaza, en TODOS los cuadros de datos de los 3 libros (excepto
Cuadro 2, que ya es la lista plana de municipios excluidos), el "0" de la
fila de un municipio excluido por el código correspondiente ("n.d."/"o.c.",
según `Excluidos/Municipios_Sin_Información.xlsx`) en TODAS las columnas de
datos de ESE bloque de ciclo (Primer o Segundo ciclo) - a pedido del
usuario 2026-10-16, para igualar la convención del libro oficial publicado
2024. Los subtotales de departamento/nacional NO se tocan (ya excluían a
estos municipios del cálculo, siguen siendo números).

Para saber qué rango de columnas pertenece a qué ciclo, se reutiliza la
MISMA fuente que ya usa `ensamblador.ensamblar()`: el pendiente cacheado de
cada paso (`hoja`, `fila_nacional`, `columna_inicio`, cantidad de
`value_cols`) - el ciclo se determina LEYENDO el texto real del encabezado
de cada bloque ("Primer ciclo nacional de vacunación..."/"Segundo
ciclo..."), no adivinando por el nombre del paso (hay cuadros "Solo Ciclo
2" cuyo paso no termina en "_c2", ej. Cuadro 5 de inventario, Cuadro 15-19
de predio-ganadero).

IMPORTANTE: este parche debe volver a correrse cada vez que se reensamble
un libro (`main.py --ensamblar <libro>`) - el reensamblado reconstruye las
hojas desde cero a partir de los pendientes calculados, que no conocen
estos códigos, así que las filas excluidas vuelven a mostrar "0" hasta que
se corre este script de nuevo.

Uso: `python -m scripts.cuadros_ganaderia.aplicar_codigos_exclusion`
(sin argumentos - aplica a los 6 archivos: `output/*_generado.xlsx` y
`Cuadros para publicación/*.xlsx` de los 3 libros)."""
from __future__ import annotations

from pathlib import Path

import openpyxl

from . import config, ensamblador, main as m

RUTA_EXCLUSION = config.BASE_DIR / "Excluidos" / "Municipios_Sin_Información.xlsx"

COL_C_COD_MUNICIPIO = 3

# Pasos que NO siguen la convención simple "1 bloque = 1 ciclo de 2025" - se
# tratan aparte:
#  - "cuadro1_ganadero_nuevos" (K-N del Cuadro 1, los 3 libros): mide
#    entidades NUEVAS en Ciclo 2 respecto a Ciclo 1 (no tiene un encabezado
#    "Primer/Segundo ciclo" propio que `_ciclo_por_encabezado` pueda leer) -
#    se usa el criterio de exclusión de Ciclo 2 (si el municipio no tiene
#    datos confiables de C2, tampoco los tiene para "nuevos").
#  - "cuadro4_ganadero"/"cuadro5_ganadero" (histórico 2023-2025, ver
#    `cuadros/ganadero_historico.py`): cada bloque compara un par de ciclos
#    DISTINTO (ej. 2024-C1 vs 2023-C2) que esta planilla de exclusión no
#    cubre (solo cubre 2025-C1/C2) - se EXCLUYEN de este parche, quedan
#    pendientes de un tratamiento propio si hace falta.
PASOS_EXCLUIR = {"cuadro4_ganadero", "cuadro5_ganadero"}
PASOS_CICLO_FORZADO = {"cuadro1_ganadero_nuevos": "C2"}

# Cuadro 12 de predio-ganadero: el usuario ajustó manualmente la fila
# "Total Nacional" de 12 a 9 en el archivo ya generado (agregó una columna
# al bloque Segundo ciclo, 2026-10-16) - ver `validacion_estructura_cuadros.py`,
# mismo ajuste reflejado acá para no desalinear el parche.
FILA_NACIONAL_OVERRIDE = {
    "cuadro12_predio_ganadero": 9,
    "cuadro12_predio_ganadero_c2": 9,
}


def _cargar_mapa_exclusion() -> dict[str, dict[str, str]]:
    """`{codigo_municipio: {"C1": codigo, "C2": codigo}}` - solo incluye el
    ciclo donde ese municipio realmente está excluido (ver columnas
    "Observación" de `RUTA_EXCLUSION`)."""
    wb = openpyxl.load_workbook(RUTA_EXCLUSION, data_only=True)
    ws = wb["Cuadro 2"]
    mapa: dict[str, dict[str, str]] = {}
    for r in range(11, ws.max_row + 1):
        cod_mun = ws.cell(row=r, column=3).value
        if not cod_mun:
            continue
        cod_mun = str(cod_mun).strip()
        obs_c1 = ws.cell(row=r, column=5).value
        obs_c2 = ws.cell(row=r, column=6).value
        entrada = {}
        if obs_c1:
            entrada["C1"] = str(obs_c1).strip()
        if obs_c2:
            entrada["C2"] = str(obs_c2).strip()
        if entrada:
            mapa[cod_mun] = entrada
    return mapa


def _ciclo_por_encabezado(ws, fila_nacional: int, col_inicio: int, col_fin: int) -> str | None:
    """Determina el ciclo de un bloque de columnas LEYENDO el texto real del
    encabezado de la hoja ("Primer ciclo nacional de vacunación..."/"Segundo
    ciclo..."). Busca, en las filas ANTES de `fila_nacional`, una celda cuyo
    rango fusionado se solape con [col_inicio, col_fin] y cuyo texto
    mencione exactamente UNO de los 2 ciclos (ignora la línea de título
    combinada "Primer Y Segundo ciclo...", que menciona ambos a la vez)."""
    rangos_por_celda = {}
    for m_ in ws.merged_cells.ranges:
        if m_.min_row < fila_nacional:
            rangos_por_celda[(m_.min_row, m_.min_col)] = m_

    for r in range(1, fila_nacional):
        for c in range(1, col_fin + 2):
            valor = ws.cell(row=r, column=c).value
            if not valor or "ciclo" not in str(valor).lower():
                continue
            texto = str(valor).lower()
            if "primer y segundo" in texto:
                continue  # línea de título combinada ("Primer y Segundo ciclo
                # nacional de vacunación...") - contiene "segundo ciclo" como
                # subcadena literal pero NO "primer ciclo" (la "y" rompe la
                # contigüidad), así que el chequeo de abajo no basta solo.
            tiene_primero = "primer ciclo" in texto
            tiene_segundo = "segundo ciclo" in texto
            if tiene_primero == tiene_segundo:
                continue  # ambos o ninguno - no sirve para decidir
            rango = rangos_por_celda.get((r, c))
            c_min, c_max = (rango.min_col, rango.max_col) if rango else (c, c)
            if c_min <= col_fin and c_max >= col_inicio:
                return "C1" if tiene_primero else "C2"
    return None


def _bloques_por_libro(nombre_libro: str, wb_referencia) -> list[tuple[str, int, int, int, str]]:
    """Devuelve [(hoja, fila_nacional, columna_inicio, columna_fin, ciclo), ...]
    para cada paso tipo "cuadro" de ese libro (excepto Cuadro 2 y los pasos
    en `PASOS_EXCLUIR`). `wb_referencia`: libro ya ensamblado (con datos y
    encabezados reales) de donde se lee el texto real del encabezado para
    inferir el ciclo de cada bloque."""
    _, _, pasos = m.LIBROS[nombre_libro]
    bloques = []
    vistos = set()
    for paso in pasos:
        if paso in PASOS_EXCLUIR or paso in vistos:
            continue
        if not ensamblador.pendiente_disponible(paso):
            continue
        _, specs = ensamblador._cargar_pendiente(paso)
        if specs.get("tipo") != "cuadro":
            continue
        hoja = specs["hoja"]
        if hoja == "Cuadro 2":
            continue
        fila_nacional = FILA_NACIONAL_OVERRIDE.get(paso, specs["fila_nacional"])
        columna_inicio = specs.get("columna_inicio", 5)
        columna_fin = columna_inicio + len(specs["value_cols"]) - 1
        ciclo = PASOS_CICLO_FORZADO.get(paso)
        if ciclo is None:
            ciclo = _ciclo_por_encabezado(wb_referencia[hoja], fila_nacional, columna_inicio, columna_fin)
        if ciclo is None:
            raise ValueError(f"No se pudo determinar el ciclo de {paso!r} ({hoja!r}, cols {columna_inicio}-{columna_fin})")
        bloques.append((hoja, fila_nacional, columna_inicio, columna_fin, ciclo))
        vistos.add(paso)
    return bloques


def _aplicar_en_archivo(ruta: Path, bloques: list[tuple[str, int, int, int, str]], mapa_exclusion: dict) -> int:
    wb = openpyxl.load_workbook(ruta)
    total_celdas = 0
    for hoja, fila_nacional, col_inicio, col_fin, ciclo in bloques:
        if hoja not in wb.sheetnames:
            continue
        ws = wb[hoja]
        for r in range(fila_nacional + 1, ws.max_row + 1):
            cod_mun = ws.cell(row=r, column=COL_C_COD_MUNICIPIO).value
            if not cod_mun:
                continue  # fila de departamento (o fin de datos) - no tocar
            cod_mun = str(cod_mun).strip()
            codigo = mapa_exclusion.get(cod_mun, {}).get(ciclo)
            if not codigo:
                continue
            for c in range(col_inicio, col_fin + 1):
                ws.cell(row=r, column=c, value=codigo)
                total_celdas += 1
    wb.save(ruta)
    return total_celdas


def main() -> None:
    mapa_exclusion = _cargar_mapa_exclusion()
    print(f"Municipios con código de exclusión: {len(mapa_exclusion)}")

    archivos_por_libro = {
        "inventario": [
            config.OUTPUT_DIR / "Cuadros caracterización inventario Ciclos 1 y 2_2025_generado.xlsx",
            config.TEMPLATE_INVENTARIO,
        ],
        "ganadero": [
            config.OUTPUT_DIR / "Cuadros caracterización del ganadero Ciclos 1 y 2_2025_generado.xlsx",
            config.TEMPLATE_GANADERO,
        ],
        "predio_ganadero": [
            config.OUTPUT_DIR / "Cuadros Caracterización predio-ganadero Ciclos 1 y 2_2025_generado.xlsx",
            config.TEMPLATE_PREDIO_GANADERO,
        ],
    }

    for libro, rutas in archivos_por_libro.items():
        # Se usa la 1ra ruta (output/..._generado.xlsx, siempre con datos y
        # encabezados reales) como referencia de texto para inferir el
        # ciclo de cada bloque - más confiable que la plantilla, que puede
        # estar vacía si nunca se copiaron resultados encima de ella.
        wb_referencia = openpyxl.load_workbook(rutas[0], data_only=True)
        bloques = _bloques_por_libro(libro, wb_referencia)
        print(f"\n{libro}: {len(bloques)} bloques detectados")
        for ruta in rutas:
            if not Path(ruta).exists():
                print(f"  (no existe, se omite: {ruta})")
                continue
            n = _aplicar_en_archivo(Path(ruta), bloques, mapa_exclusion)
            print(f"  {ruta.name}: {n} celdas reemplazadas")


if __name__ == "__main__":
    main()
