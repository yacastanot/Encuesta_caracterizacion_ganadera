"""Cuadro 4 y 5 del libro "ganadero" - "Cantidad de ganaderos, de nuevos
ganaderos y de ganaderos que se mantienen" respecto al ciclo anterior
(Cuadro 4) / al mismo ciclo del año anterior (Cuadro 5).

Necesitan 6 "fotos" de identidad de ganaderos (`IDENT_GANADERO` por
municipio, con su `peso_ganadero`): 2023-C1, 2023-C2, 2024-C1, 2024-C2 (ver
`base_maestra_ganadero_historico.py` - 2024 reconstruido desde cero, 2023
tomado tal cual del archivo más reciente del pipeline legado, decisión del
usuario) y 2025-C1/2025-C2 (ya existentes, `base_maestra_c1.py`/`_c2.py`).

Cuadro 4 (ciclo cronológico anterior) - 4 bloques, cada uno comparado contra
el ciclo INMEDIATAMENTE ANTERIOR en el calendario de vacunación:
    2024-C1 vs 2023-C2  |  2024-C2 vs 2024-C1  |  2025-C1 vs 2024-C2  |  2025-C2 vs 2025-C1
Cuadro 5 (mismo ciclo, año anterior) - 4 bloques, cada uno comparado contra
el MISMO ciclo del año anterior:
    2024-C1 vs 2023-C1  |  2025-C1 vs 2024-C1  |  2024-C2 vs 2023-C2  |  2025-C2 vs 2024-C2

Para cada bloque (ciclo "actual" vs. ciclo "referencia"), por municipio:
 - "Total de ganaderos en el ciclo" = SUM(peso_ganadero) del ciclo actual.
 - "nuevos respecto al ciclo/año anterior" = SUM(peso_ganadero) del ciclo
   actual, de los `IDENT_GANADERO` que NO estaban en el ciclo de referencia.
 - "salieron respecto al ciclo/año anterior" = SUM(peso_ganadero) del ciclo
   de REFERENCIA (no del actual - un ganadero que salió no tiene fila en el
   ciclo actual), de los `IDENT_GANADERO` que NO están en el ciclo actual -
   se agrupa por el municipio que tenían EN LA REFERENCIA (es donde estaban,
   no hay forma de saber dónde estarían ahora).
 - "se mantienen en los dos ciclos" = SUM(peso_ganadero) del ciclo actual, de
   los `IDENT_GANADERO` que SÍ estaban en la referencia.

"nuevos" + "se mantienen" = "Total" es una partición exacta del ciclo actual
(todo ganadero actual es nuevo O se mantiene, nunca ambos, "salieron" es una
cantidad del ciclo de REFERENCIA que no forma parte del total actual) - SÍ
se declara como `columnas_totales` (corregido 2026-09-25, el usuario detectó
el residuo de redondeo independiente: una versión anterior dejaba "total"
independiente "para coincidir con Cuadro 1/3", pero esa razón no aplicaba -
mismo criterio ya corregido en Cuadro 7, ver `cuadros/ganadero.py`). Esto
aplica a los bloques calculados (2025-C1/2025-C2); los bloques 2024
copiados del libro publicado (ver más abajo) usan el "Total" tal cual viene
publicado, sin recalcularlo de sus componentes.

TRADEOFF ACEPTADO (2026-09-25, decisión del usuario): al derivar "total" de
nuevos+mantienen en Cuadro 4 y Cuadro 5 por separado, el "total_ganaderos"
de 2025-C1/2025-C2 YA NO coincide exactamente entre Cuadro 4 y Cuadro 5 -
antes sí, porque "total" se redondeaba independiente de la partición usada.
La causa: Cuadro 4 usa el ciclo INMEDIATAMENTE anterior como referencia y
Cuadro 5 usa el mismo ciclo del año anterior - son 2 particiones distintas
de la misma población 2025, y el redondeo independiente por municipio de
cada partición no da exactamente la misma suma. Residuo verificado: ~0.005%
a nivel nacional (32/682.243 en 2025-C1, 12/644.461 en 2025-C2). El usuario
priorizó la identidad interna de cada cuadro (Total=Nuevos+Mantienen, fila
por fila) sobre la coincidencia cruzada Cuadro4↔Cuadro5 - por eso
`validacion_estructura_cuadros.py` ya NO declara esa comparación cruzada
(antes sí la tenía, cuando aplicaba).

CORRECCIÓN 2026-09-25 (decisión del usuario) - LOS BLOQUES 2024 (2024-C1,
2024-C2) SE MUESTRAN CON LOS VALORES YA PUBLICADOS, NO CON EL CÁLCULO
INDEPENDIENTE DE ARRIBA: validado municipio por municipio contra
`anex-CAG-CaractGanadero-2024.xlsx` (el libro ya publicado), el cálculo
propio daba diferencias reales en 1.036 de 1.121 municipios (hasta 4% de
diferencia nacional en "Total", hasta 30% en "salieron") - la causa más
probable es que el candidato de encuesta 2024 elegido por fecha de archivo
más reciente (`encuestac1corregida.sas7bdat`) no es necesariamente el mismo
snapshot que usó el equipo que publicó el libro 2024 (había ~10 versiones
candidatas sin forma de saber cuál es "la" oficial, ver `config.py`). Ante
eso, se copian los 4 valores (total/nuevos/salieron/mantienen) de las
columnas correspondientes de Cuadro 4 (M-P="1er Ciclo 2024", Q-T="2do Ciclo
2024") y Cuadro 5 (I-L="1er Ciclo 2024", Q-T="2do Ciclo 2024") del libro
2024, a nivel MUNICIPIO, redondeados a entero, con departamento/nacional
recalculados como suma de esos municipios YA redondeados (`agregador.generar_cuadro`)
- así la consistencia en los 3 niveles queda garantizada aunque el libro 2024
publicado no lo estuviera (sus filas Nacional/Departamento guardan el valor
decimal sin redondear, no necesariamente la suma exacta de sus municipios).
Municipios marcados "n.d."/"o.c."/"ZLSV" en el libro 2024 (sin dato) se
tratan como 0.

IMPORTANTE: el cálculo independiente de 2024-C1/2024-C2 (arriba) SIGUE
haciendo falta y NO se elimina - los bloques "2025-C1 vs 2024-C2"
(Cuadro 4) y "2025-C1 vs 2024-C1" (Cuadro 5) necesitan la identidad
`IDENT_GANADERO` de 2024 a nivel predio-ganadero (no solo su agregado) para
calcular CUÁLES ganaderos de 2025 son nuevos/se mantienen respecto a 2024 -
el libro publicado 2024 no trae esa identidad, solo agregados. Por eso acá
se calcula TODO primero (como antes) y solo al final se SOBRESCRIBEN las
columnas `c2024c1_*`/`c2024c2_*` con los valores publicados - las columnas
`c2025c1_*`/`c2025c2_*` quedan con el cálculo propio, sin tocar.
"""
from __future__ import annotations

import openpyxl
import pandas as pd

from .. import agregador, base_maestra_c1, base_maestra_c2, base_maestra_ganadero_historico, config

_COLS = ["predioganid", "IDENT_GANADERO", "CODIGO_MUNICIPIO", "peso_ganadero"]

RUTA_PUBLICADO_2024 = config.BASE_DIR / "01Entrada" / "Cuadros publicados 2024" / "anex-CAG-CaractGanadero-2024.xlsx"

# `(hoja, columna Total)` del libro 2024 - las otras 3 columnas del bloque
# (nuevos/salieron/mantienen) son las 3 siguientes, en ese orden, verificado
# contra los encabezados reales (fila 10) de cada hoja.
_BLOQUES_PUBLICADOS = {
    ("cuadro4", "c2024c1"): ("Cuadro 4", 13),  # M = "1er Ciclo Nacional de Vacunación 2024"
    ("cuadro4", "c2024c2"): ("Cuadro 4", 17),  # Q = "2do Ciclo Nacional de Vacunación 2024"
    ("cuadro5", "c2024c1"): ("Cuadro 5", 9),   # I = "1er Ciclo Nacional de Vacunación 2024"
    ("cuadro5", "c2024c2"): ("Cuadro 5", 17),  # Q = "2do Ciclo Nacional de Vacunación 2024"
}


def _valor_numerico(v) -> float:
    """Las celdas de municipios sin dato para ese ciclo traen texto ("n.d.",
    "o.c.", "ZLSV") en vez de un número - se tratan como 0 (mismo criterio
    que "sin registro" en el resto del proyecto)."""
    return float(v) if isinstance(v, (int, float)) else 0.0


def _bloque_publicado_2024(cuadro: str, prefijo: str) -> pd.DataFrame:
    """Lee, a nivel MUNICIPIO, el bloque (total/nuevos/salieron/mantienen) de
    `anex-CAG-CaractGanadero-2024.xlsx` que corresponde a `(cuadro, prefijo)`
    - ver `_BLOQUES_PUBLICADOS` y el docstring del módulo."""
    hoja, col_total = _BLOQUES_PUBLICADOS[(cuadro, prefijo)]
    wb = openpyxl.load_workbook(RUTA_PUBLICADO_2024, data_only=True)
    ws = wb[hoja]

    filas = []
    for r in range(11, ws.max_row + 1):
        codigo = ws.cell(row=r, column=3).value
        if not codigo or len(str(codigo)) != 5:
            continue
        filas.append({
            "CODIGO_MUNICIPIO": str(codigo).zfill(5),
            f"{prefijo}_total": round(_valor_numerico(ws.cell(row=r, column=col_total).value)),
            f"{prefijo}_nuevos": round(_valor_numerico(ws.cell(row=r, column=col_total + 1).value)),
            f"{prefijo}_salieron": round(_valor_numerico(ws.cell(row=r, column=col_total + 2).value)),
            f"{prefijo}_mantienen": round(_valor_numerico(ws.cell(row=r, column=col_total + 3).value)),
        })
    municipios = pd.DataFrame(filas)
    cols = [f"{prefijo}_total", f"{prefijo}_nuevos", f"{prefijo}_salieron", f"{prefijo}_mantienen"]
    # Reagrega Departamento/Nacional como suma de los municipios YA
    # redondeados (garantiza consistencia en los 3 niveles, sin depender de
    # si las filas Nacional/Departamento del libro 2024 ya lo eran).
    return agregador.generar_cuadro(municipios, cols)


def _cargar(etiqueta: str) -> pd.DataFrame:
    if etiqueta == "2025-C1":
        df = pd.read_parquet(base_maestra_c1.RUTA_BASE_MAESTRA_C1, columns=_COLS)
    elif etiqueta == "2025-C2":
        df = pd.read_parquet(base_maestra_c2.RUTA_BASE_MAESTRA_C2, columns=_COLS)
    elif etiqueta == "2024-C1":
        df = base_maestra_ganadero_historico.construir_2024(1)
    elif etiqueta == "2024-C2":
        df = base_maestra_ganadero_historico.construir_2024(2)
    elif etiqueta == "2023-C1":
        df = base_maestra_ganadero_historico.leer_2023(1)
    elif etiqueta == "2023-C2":
        df = base_maestra_ganadero_historico.leer_2023(2)
    else:
        raise ValueError(f"etiqueta de ciclo desconocida: {etiqueta!r}")
    df["CODIGO_MUNICIPIO"] = df["CODIGO_MUNICIPIO"].astype(str).str.zfill(5)
    # `IDENT_GANADERO` viene en tipos incompatibles entre fuentes: `float` en
    # todo lo que sale de un `.sas7bdat` vía pyreadstat (base_maestra_c1/2025,
    # y las 4 fotos históricas de acá, TODAS de origen SAS) vs texto en
    # `base_maestra_c2`/2025 (construida desde CSV) - verificado 2026-09-24:
    # sin normalizar, el bloque "2025-C2 vs 2025-C1"/"2025-C2 vs 2024-C2" daba
    # 100% "nuevos" (0 coincidencias reales) - MISMO bug ya documentado y
    # corregido en `cuadros/comun.py` (`_ids_c1_normalizados`) para las
    # columnas "nuevos" de Cuadro 1.
    #
    # El lado texto (C2) tiene 1 solo valor sucio en las 712.846 filas
    # ("686805_", con un guion bajo pegado - verificado que es el ÚNICO caso,
    # no un patrón repetido) que rompe `pd.to_numeric` directo - se limpia
    # quitando cualquier caracter que no sea dígito ANTES de convertir, en vez
    # de dejarlo tal cual como texto sucio (que nunca haría match con el
    # "686805" limpio de otro ciclo, si existiera): es un error de captura
    # evidente, no una identidad realmente distinta.
    if pd.api.types.is_float_dtype(df["IDENT_GANADERO"]):
        ident = pd.to_numeric(df["IDENT_GANADERO"])
    else:
        ident = pd.to_numeric(df["IDENT_GANADERO"].astype(str).str.replace(r"[^0-9]", "", regex=True))
    df["IDENT_GANADERO"] = ident.astype("Int64").astype(str)
    return df


def _comparar_bloque(actual: pd.DataFrame, referencia: pd.DataFrame, prefijo: str) -> tuple[pd.DataFrame, list[str]]:
    ids_actual = set(actual["IDENT_GANADERO"])
    ids_referencia = set(referencia["IDENT_GANADERO"])

    es_nuevo = ~actual["IDENT_GANADERO"].isin(ids_referencia)
    es_salido = ~referencia["IDENT_GANADERO"].isin(ids_actual)

    total = actual.groupby("CODIGO_MUNICIPIO")["peso_ganadero"].sum().rename(f"{prefijo}_total")
    nuevos = actual[es_nuevo].groupby("CODIGO_MUNICIPIO")["peso_ganadero"].sum().rename(f"{prefijo}_nuevos")
    mantienen = actual[~es_nuevo].groupby("CODIGO_MUNICIPIO")["peso_ganadero"].sum().rename(f"{prefijo}_mantienen")
    salieron = referencia[es_salido].groupby("CODIGO_MUNICIPIO")["peso_ganadero"].sum().rename(f"{prefijo}_salieron")

    cols = [f"{prefijo}_total", f"{prefijo}_nuevos", f"{prefijo}_salieron", f"{prefijo}_mantienen"]
    agg = pd.concat([total, nuevos, salieron, mantienen], axis=1).reset_index()
    for c in cols:
        agg[c] = agg[c].fillna(0.0)
    return agg[["CODIGO_MUNICIPIO"] + cols], cols


def _armar_cuadro(nombre_cuadro: str, bloques: list[tuple[str, str, str]]) -> tuple[pd.DataFrame, list[str]]:
    """`bloques`: lista de `(prefijo, etiqueta_actual, etiqueta_referencia)`,
    en el orden en que van las columnas en la plantilla (E-H, I-L, M-P, Q-T).
    `nombre_cuadro`: "cuadro4"/"cuadro5" - qué bloques sobrescribir con datos
    ya publicados de 2024, ver `_BLOQUES_PUBLICADOS` y el docstring del
    módulo."""
    agg = None
    value_cols: list[str] = []
    columnas_totales: dict[str, list[str]] = {}
    for prefijo, actual_et, referencia_et in bloques:
        actual = _cargar(actual_et)
        referencia = _cargar(referencia_et)
        bloque_df, cols = _comparar_bloque(actual, referencia, prefijo)
        value_cols += cols
        # "total" = "nuevos" + "se mantienen" (partición exacta del ciclo
        # ACTUAL - "salieron" es una cantidad del ciclo de REFERENCIA, no
        # forma parte del total actual, ver docstring del módulo). Corregido
        # 2026-09-25 (el usuario detectó el residuo, mismo criterio que
        # Cuadro 7): antes "total" quedaba independiente "para coincidir con
        # Cuadro 1/3", pero esa razón no aplicaba - ya no se declara así.
        columnas_totales[f"{prefijo}_total"] = [f"{prefijo}_nuevos", f"{prefijo}_mantienen"]
        agg = bloque_df if agg is None else agg.merge(bloque_df, on="CODIGO_MUNICIPIO", how="outer")
    for c in value_cols:
        agg[c] = agg[c].fillna(0.0)

    tabla = agregador.generar_cuadro(agg, value_cols, columnas_totales=columnas_totales)

    # Sobrescribe los bloques 2024 con los valores YA PUBLICADOS (ver
    # docstring del módulo) - los bloques 2025 quedan con el cálculo propio.
    llaves = ["NIVEL", "COD_DEPARTAMENTO", "DEPARTAMENTO", "CODIGO_MUNICIPIO", "MUNICIPIO"]
    for prefijo, _, _ in bloques:
        if (nombre_cuadro, prefijo) not in _BLOQUES_PUBLICADOS:
            continue
        cols_bloque = [f"{prefijo}_total", f"{prefijo}_nuevos", f"{prefijo}_salieron", f"{prefijo}_mantienen"]
        publicado = _bloque_publicado_2024(nombre_cuadro, prefijo)
        tabla = tabla.drop(columns=cols_bloque).merge(publicado, on=llaves, how="left")
        tabla[cols_bloque] = tabla[cols_bloque].fillna(0)
    tabla = tabla[llaves + value_cols]

    return tabla, value_cols


def generar_cuadro4() -> tuple[pd.DataFrame, list[str]]:
    """Cuadro 4: cada bloque vs. el ciclo cronológico inmediatamente anterior."""
    return _armar_cuadro("cuadro4", [
        ("c2024c1", "2024-C1", "2023-C2"),
        ("c2024c2", "2024-C2", "2024-C1"),
        ("c2025c1", "2025-C1", "2024-C2"),
        ("c2025c2", "2025-C2", "2025-C1"),
    ])


def generar_cuadro5() -> tuple[pd.DataFrame, list[str]]:
    """Cuadro 5: cada bloque vs. el mismo ciclo del año anterior. Orden de
    bloques de la plantilla real: 2024-C1, 2025-C1, 2024-C2, 2025-C2 (NO
    2024-C1/2024-C2/2025-C1/2025-C2 - verificado contra los encabezados
    reales de la hoja)."""
    return _armar_cuadro("cuadro5", [
        ("c2024c1", "2024-C1", "2023-C1"),
        ("c2025c1", "2025-C1", "2024-C1"),
        ("c2024c2", "2024-C2", "2023-C2"),
        ("c2025c2", "2025-C2", "2024-C2"),
    ])
