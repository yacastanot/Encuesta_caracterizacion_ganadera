"""Escritura de resultados dentro de las plantillas Excel originales.

Estrategia para que el formato quede EXACTAMENTE igual al de Adriana: nunca se
construye un archivo desde cero. Se abre la plantilla real (título,
encabezados multinivel, merges, notas al pie ya existen ahí) y solo se
inserta/llena el bloque de filas de datos (nacional + departamentos +
municipios), clonando el estilo de celda que ya trae la fila de ejemplo
"Total Nacional" / "Total por departamento" de cada hoja.

IDEMPOTENCIA: como cada libro se arma cuadro por cuadro (cada uno en su propio
proceso Python, por límites de memoria de la máquina - ver `main.py`), es
normal tener que re-escribir la MISMA hoja más de una vez (ej. para corregir
un bug). Por eso los estilos de la fila "Total Nacional", de la fila de
departamento, y el contenido/estilo de las notas al pie SIEMPRE se leen de la
plantilla PRÍSTINA (`ruta_plantilla`/`wb_pristina`), nunca de lo que ya haya
en el libro que se está armando - y antes de escribir, se borra por completo
cualquier fila que una corrida anterior haya dejado en esa hoja (desde
`fila_nacional` en adelante). Así, escribir un cuadro por segunda vez da
exactamente el mismo resultado que la primera vez, en vez de ir acumulando
filas huérfanas.

DOS NIVELES DE API:
 - `_escribir_cuadro_en_wb` / `_escribir_lista_plana_en_wb`: la lógica real,
   opera sobre un `openpyxl.Workbook` YA ABIERTO en memoria, sin cargar ni
   guardar nada. Pensada para `ensamblador.py`, que carga la plantilla UNA
   sola vez, llama esto una vez por cuadro, y guarda UNA sola vez al final -
   evita que un libro que va creciendo (varios MB, muchas hojas) se recargue
   y regrabe por completo en cada cuadro (con 9 cuadros en el libro de
   inventario, eso llegó a tomar ~17 minutos en total; ver `ensamblador.py`).
 - `escribir_cuadro` / `escribir_lista_plana`: wrappers que cargan, llaman lo
   anterior y guardan - para uso puntual/aislado (debugging, un solo cuadro
   suelto). El pipeline real ya no los usa para armar los 3 libros completos.
"""
from __future__ import annotations

from copy import copy

import openpyxl
import pandas as pd
from openpyxl.styles import PatternFill
from openpyxl.styles.colors import Color
from openpyxl.worksheet.worksheet import Worksheet

from . import agregador

COL_A_COD_DEPARTAMENTO = 1
COL_B_DEPARTAMENTO = 2
COL_C_COD_MUNICIPIO = 3
COL_D_MUNICIPIO = 4
PRIMERA_COLUMNA_DATOS = 5  # columna E

# Sombreado alterno de filas de datos (departamento/municipio) - mismo color
# que usa el libro publicado 2024 (`anex-CAG-Caract*-2024.xlsx`, verificado
# en las 3 hojas "Cuadro 1": relleno gris muy claro tema0/tint≈-0.05 en la
# fila de índice PAR dentro del bloque departamento+municipio - fila 12 en
# 2024, la primera fila de datos - y SIN relleno en la fila de índice impar,
# alternando de ahí en adelante fila por fila, no por departamento).
# Decisión del usuario (2026-09-24): igualar el formato del libro 2024,
# aplicado acá en código porque la plantilla clona UN solo estilo para todas
# las filas de datos (ver `_escribir_cuadro_en_wb`) - alternar el relleno no
# se puede lograr solo editando la plantilla.
_RELLENO_FILA_PAR = PatternFill(patternType="solid", fgColor=Color(theme=0, tint=-0.0499893185216834))
_SIN_RELLENO = PatternFill(fill_type=None)


def _clonar_estilo_celda(origen, destino) -> None:
    destino.font = copy(origen.font)
    destino.border = copy(origen.border)
    destino.fill = copy(origen.fill)
    destino.number_format = origen.number_format
    destino.alignment = copy(origen.alignment)
    destino.protection = copy(origen.protection)


def _clonar_estilo_fila(ws_origen: Worksheet, fila_origen: int, ws_destino: Worksheet, fila_destino: int, n_columnas: int) -> None:
    for c in range(1, n_columnas + 1):
        _clonar_estilo_celda(ws_origen.cell(row=fila_origen, column=c), ws_destino.cell(row=fila_destino, column=c))


def _aplicar_sombreado_alterno(ws_destino: Worksheet, fila_destino: int, n_columnas: int, indice: int, columna_inicio: int = 1) -> None:
    """Sobreescribe SOLO el relleno (no fuente/borde, ya clonados aparte) de
    una fila de datos según la paridad de `indice` (0-based dentro del
    bloque departamento+municipio) - ver constantes arriba."""
    relleno = _RELLENO_FILA_PAR if indice % 2 == 0 else _SIN_RELLENO
    for c in range(columna_inicio, n_columnas + 1):
        ws_destino.cell(row=fila_destino, column=c).fill = copy(relleno)


def _desfusionar_desde(ws: Worksheet, fila_desde: int) -> None:
    """Elimina cualquier celda combinada de `ws` que empiece en o después de
    `fila_desde` - defensivo antes de borrar/escribir esas filas.

    Por qué: `ws.delete_rows()` de openpyxl NO limpia las combinaciones de
    celdas que caen dentro del rango borrado (bug conocido de la librería) -
    la combinación "fantasma" sobrevive en el archivo guardado. Como una
    celda combinada solo conserva el valor de su celda superior-izquierda,
    esto hace que Excel (y openpyxl al releer) descarte silenciosamente los
    valores de las demás columnas de esa fila - se detectó así: la plantilla
    de "Cuadro 2" trae una nota al pie combinada en A27:F27 (pensada para
    cuando la lista de municipios excluidos tenía ~16 filas), y al crecer esa
    lista a 67, los datos de las filas 27+ se escribían bien en memoria pero
    se veían truncados a 16 filas tras guardar y releer el archivo."""
    afectadas = [str(m) for m in ws.merged_cells.ranges if m.min_row >= fila_desde]
    for rango in afectadas:
        ws.unmerge_cells(rango)


def _buscar_fila_notas(ws: Worksheet, fila_desde: int, limite: int = 20000) -> int | None:
    """Busca la primera fila con contenido real en o después de `fila_desde`
    en la hoja PRÍSTINA (nunca en una hoja que ya hayamos escrito nosotros)."""
    for r in range(fila_desde, min(ws.max_row, limite) + 1):
        if any(ws.cell(row=r, column=c).value not in (None, "") for c in range(1, ws.max_column + 1)):
            return r
    return None


def _capturar_bloque_notas(ws_pristina: Worksheet, fila_desde: int, n_columnas: int):
    """Devuelve la lista de filas de notas al pie de la plantilla PRÍSTINA
    (valores + estilos), como bloque contiguo a partir de la primera fila con
    contenido encontrada desde `fila_desde`. None si no hay ninguna."""
    fila_inicio = _buscar_fila_notas(ws_pristina, fila_desde)
    if fila_inicio is None:
        return None
    filas = []
    r = fila_inicio
    vacio_consecutivo = 0
    while r <= ws_pristina.max_row and vacio_consecutivo < 3:
        vals = [ws_pristina.cell(row=r, column=c).value for c in range(1, max(n_columnas, ws_pristina.max_column) + 1)]
        if any(v not in (None, "") for v in vals):
            filas.append(r)
            vacio_consecutivo = 0
        else:
            vacio_consecutivo += 1
        r += 1
    return filas


def _escribir_cuadro_en_wb(
    wb: openpyxl.Workbook,
    wb_pristina: openpyxl.Workbook,
    hoja: str,
    fila_nacional: int,
    tabla: pd.DataFrame,
    value_cols: list[str],
    formato_numero: str = "#,##0",
    columna_inicio: int = PRIMERA_COLUMNA_DATOS,
    encabezados: dict[str, str] | None = None,
    nota_extra_etiqueta: str | None = None,
    nota_extra_texto: str | None = None,
) -> None:
    """Inserta `tabla` (salida de `agregador.generar_cuadro`) dentro de `hoja`,
    a partir de la columna `columna_inicio` (por defecto E, la del bloque
    "Primer ciclo"; usar la columna real del bloque "Segundo ciclo" - ej. AA=27
    en Cuadro 3/7, T=20 en Cuadro 8 - para llenar ese otro bloque de la MISMA
    hoja sin tocar el ya escrito).

    `fila_nacional`: número de fila donde va la fila "Total Nacional" (según
    la plantilla PRÍSTINA original - exista ya ese rótulo ahí o no).

    `encabezados`: opcional, `{celda: texto}` (ej. `{"F6": "0 a 2
    trabajadores"}`) para SOBRESCRIBIR encabezados de la plantilla que quedan
    incompletos o son placeholders (ej. Cuadro 5 de predio-ganadero, que trae
    literalmente "XXX hasta XXXX" a la espera de que se definan los rangos) -
    únicamente estas celdas puntuales, nunca se toca el resto del encabezado.

    `nota_extra_etiqueta`/`nota_extra_texto`: opcional, agrega UNA fila más de
    "nota al pie" (mismo estilo/posición que las que ya trae la plantilla,
    justo después de ellas) documentando una decisión metodológica que NO
    viene de la plantilla original (ej. cómo se definieron unos rangos que la
    plantilla dejaba sin resolver) - sin modificar el archivo de plantilla en
    disco, queda documentado en el código (acá) y se reescribe igual en cada
    corrida.

    Si `columna_inicio` es la del "Primer ciclo" (el valor por defecto), la
    hoja se reconstruye desde cero: se borra todo lo que haya desde
    `fila_nacional` en adelante y se vuelve a armar el esqueleto (filas +
    notas al pie), así que es segura de llamar más de una vez sobre el mismo
    `wb` para ESE bloque. Si es la de un bloque adicional (ej. "Segundo
    ciclo") sobre una hoja que ya tiene su esqueleto escrito (por una llamada
    anterior sobre el MISMO `wb`, ver `ensamblador.py` para el orden), NO se
    borran filas: solo se clona el estilo y se llenan los valores de esas
    columnas nuevas, dejando intacto lo que ya había.

    No carga ni guarda nada - `wb`/`wb_pristina` ya deben estar abiertos (ver
    docstring del módulo). `wb_pristina` puede ser el mismo objeto para todas
    las llamadas de un mismo ensamblado: solo se lee de ahí, nunca se muta.
    """
    ws = wb[hoja]
    n_columnas = columna_inicio + len(value_cols) - 1
    ws_pristina = wb_pristina[hoja]

    if encabezados:
        for celda, texto in encabezados.items():
            ws[celda] = texto

    nacional = tabla[tabla["NIVEL"] == agregador.NIVEL_NACIONAL].reset_index(drop=True)
    depmun = tabla[tabla["NIVEL"] != agregador.NIVEL_NACIONAL].reset_index(drop=True)
    filas_nacional_usadas = max(len(nacional), 1)
    fila_dept_inicio = fila_nacional + filas_nacional_usadas

    bloque_adicional = columna_inicio != PRIMERA_COLUMNA_DATOS and ws.max_row >= fila_dept_inicio

    if not bloque_adicional:
        # La convención general es que `fila_nacional + 1` es una fila de
        # ESTILO (ej. "Total por departamento", columna A vacía) y las notas
        # al pie empiezan más abajo, desde `fila_nacional + 2`. Algunas
        # plantillas (ej. Cuadro 8 del libro "ganadero") no traen esa fila de
        # estilo y ponen una nota (columna A con texto, ej. "Pregunta del
        # RUV") justo en `fila_nacional + 1` - si no se detecta esto, esa
        # nota se trata como si fuera la fila de estilo departamental y se
        # pierde por completo al borrar/reescribir. Se detecta por columna A
        # no vacía (una fila de estilo real siempre la deja en blanco).
        hay_nota_pegada = bool(ws_pristina.cell(row=fila_nacional + 1, column=COL_A_COD_DEPARTAMENTO).value)
        fila_notas_desde = fila_nacional + 1 if hay_nota_pegada else fila_nacional + 2
        filas_notas_pristina = _capturar_bloque_notas(ws_pristina, fila_notas_desde, n_columnas)

        # --- Limpieza idempotente: borrar cualquier resto de una corrida anterior ---
        if ws.max_row >= fila_nacional:
            _desfusionar_desde(ws, fila_nacional)
            ws.delete_rows(fila_nacional, amount=ws.max_row - fila_nacional + 1)

        # --- Bloque nacional (puede ser 1 fila, o varias si el cuadro desagrega
        # por una categoría extra a nivel nacional). Estilo SIEMPRE desde la
        # plantilla prístina, fila `fila_nacional`. ---
        for i in range(filas_nacional_usadas):
            f = fila_nacional + i
            _clonar_estilo_fila(ws_pristina, fila_nacional, ws, f, n_columnas)
            ws.cell(row=f, column=COL_A_COD_DEPARTAMENTO, value="")
            ws.cell(row=f, column=COL_B_DEPARTAMENTO, value="Total Nacional")

        # --- Notas al pie: se vuelven a pegar tal cual estaban en la plantilla
        # prístina, inmediatamente después de donde va a terminar el bloque de
        # datos (más abajo se inserta el espacio que necesitan los datos). Si
        # se pidió `nota_extra_texto`, se agrega UNA fila más al final de ese
        # mismo bloque (ver docstring de la función). ---
        n_necesarias = len(depmun)
        n_notas_pristina = len(filas_notas_pristina) if filas_notas_pristina else 0
        if filas_notas_pristina:
            for offset, r_pristina in enumerate(filas_notas_pristina):
                f_destino = fila_dept_inicio + offset
                for c in range(1, max(n_columnas, ws_pristina.max_column) + 1):
                    src = ws_pristina.cell(row=r_pristina, column=c)
                    dst = ws.cell(row=f_destino, column=c, value=src.value)
                    _clonar_estilo_celda(src, dst)
        if nota_extra_texto:
            f_destino = fila_dept_inicio + n_notas_pristina
            estilo_origen = filas_notas_pristina[0] if filas_notas_pristina else fila_nacional + 1
            _clonar_estilo_fila(ws_pristina, estilo_origen, ws, f_destino, n_columnas)
            if nota_extra_etiqueta:
                ws.cell(row=f_destino, column=COL_B_DEPARTAMENTO, value=nota_extra_etiqueta)
                ws.cell(row=f_destino, column=COL_C_COD_MUNICIPIO, value=nota_extra_texto)
            else:
                ws.cell(row=f_destino, column=COL_B_DEPARTAMENTO, value=nota_extra_texto)
        if filas_notas_pristina or nota_extra_texto:
            ws.insert_rows(fila_dept_inicio, amount=n_necesarias)

        # --- Bloque departamentos + municipios. Estilo SIEMPRE desde la
        # plantilla prístina, fila `fila_nacional + 1` (rótulo original "Total
        # por departamento" o la primera fila vacía si ese placeholder no
        # existía) - o `fila_nacional` mismo si esa fila resultó ser una nota
        # pegada (ver arriba), reutilizando el estilo de la fila nacional. ---
        estilo_depto_origen = fila_nacional if hay_nota_pegada else fila_nacional + 1
        for i, row in depmun.iterrows():
            f = fila_dept_inicio + i
            _clonar_estilo_fila(ws_pristina, estilo_depto_origen, ws, f, n_columnas)
            _aplicar_sombreado_alterno(ws, f, n_columnas, i)
            ws.cell(row=f, column=COL_A_COD_DEPARTAMENTO, value=row["COD_DEPARTAMENTO"])
            ws.cell(row=f, column=COL_B_DEPARTAMENTO, value=row["DEPARTAMENTO"])
            if row["NIVEL"] == agregador.NIVEL_MUNICIPIO:
                ws.cell(row=f, column=COL_C_COD_MUNICIPIO, value=row["CODIGO_MUNICIPIO"])
                ws.cell(row=f, column=COL_D_MUNICIPIO, value=row["MUNICIPIO"])
            else:
                ws.cell(row=f, column=COL_C_COD_MUNICIPIO, value=None)
                ws.cell(row=f, column=COL_D_MUNICIPIO, value=None)
    else:
        # Bloque adicional (ej. "Segundo ciclo") sobre una hoja cuyo esqueleto
        # (filas + notas) ya escribió el bloque "Primer ciclo": solo se clona
        # el estilo de ESTAS columnas nuevas, sin borrar/mover nada existente.
        for i in range(filas_nacional_usadas):
            f = fila_nacional + i
            for c in range(columna_inicio, n_columnas + 1):
                _clonar_estilo_celda(ws_pristina.cell(row=fila_nacional, column=c), ws.cell(row=f, column=c))
        for i in range(len(depmun)):
            f = fila_dept_inicio + i
            for c in range(columna_inicio, n_columnas + 1):
                _clonar_estilo_celda(ws_pristina.cell(row=fila_nacional + 1, column=c), ws.cell(row=f, column=c))
            _aplicar_sombreado_alterno(ws, f, n_columnas, i, columna_inicio=columna_inicio)

    # Los cuadros publicados muestran cantidades enteras (predios, ganaderos,
    # animales) - `formato_numero` ("#,##0") solo controla cómo se VE la
    # celda en Excel, el valor guardado seguía siendo el decimal completo
    # (ej. 524.698735, de los pesos de calibración). Se redondea acá, en el
    # único punto donde se escribe el valor final, para que el valor
    # guardado sea entero de verdad - los cálculos/validaciones previos
    # siguen usando la precisión completa, esto solo afecta la presentación.
    for i, row in nacional.iterrows():
        f = fila_nacional + i
        for j, col in enumerate(value_cols):
            celda = ws.cell(row=f, column=columna_inicio + j, value=round(float(row[col])))
            celda.number_format = formato_numero

    for i, row in depmun.iterrows():
        f = fila_dept_inicio + i
        for j, col in enumerate(value_cols):
            celda = ws.cell(row=f, column=columna_inicio + j, value=round(float(row[col])))
            celda.number_format = formato_numero


def escribir_cuadro(
    ruta_plantilla,
    ruta_salida,
    hoja: str,
    fila_nacional: int,
    tabla: pd.DataFrame,
    value_cols: list[str],
    formato_numero: str = "#,##0",
    columna_inicio: int = PRIMERA_COLUMNA_DATOS,
) -> None:
    """Wrapper de un solo uso: carga, escribe UN cuadro y guarda. Para armar
    un libro completo con varios cuadros, usar `ensamblador.py` (carga una
    vez, escribe todos los cuadros con `_escribir_cuadro_en_wb`, guarda una
    vez) - ver docstring del módulo para el porqué."""
    origen = ruta_salida if ruta_salida.exists() else ruta_plantilla
    wb = openpyxl.load_workbook(origen)
    wb_pristina = openpyxl.load_workbook(ruta_plantilla)
    _escribir_cuadro_en_wb(wb, wb_pristina, hoja, fila_nacional, tabla, value_cols, formato_numero, columna_inicio)
    ruta_salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta_salida)


def _escribir_lista_plana_en_wb(
    wb: openpyxl.Workbook,
    hoja: str,
    fila_inicio: int,
    tabla: pd.DataFrame,
    columnas: list[str],
    columna_inicio: int = COL_A_COD_DEPARTAMENTO,
) -> None:
    """Escribe `tabla` como lista plana (sin jerarquía Nacional/Departamento/
    Municipio), a partir de `fila_inicio`, una columna de `tabla` por columna
    de la hoja. Para cuadros como el "Cuadro 2 - Municipios sin información
    asociada", que no pasan por `agregador.generar_cuadro`. Idempotente igual
    que `_escribir_cuadro_en_wb`: borra cualquier resto de una corrida
    anterior antes de escribir. Sin plantilla de estilo de fila (la plantilla
    original no trae ninguna fila de ejemplo con datos para este cuadro) - se
    escribe con el estilo por defecto de openpyxl.

    No carga ni guarda nada - `wb` ya debe estar abierto (ver docstring del
    módulo)."""
    ws = wb[hoja]

    if ws.max_row >= fila_inicio:
        _desfusionar_desde(ws, fila_inicio)
        ws.delete_rows(fila_inicio, amount=ws.max_row - fila_inicio + 1)

    for i, row in tabla.reset_index(drop=True).iterrows():
        f = fila_inicio + i
        for j, col in enumerate(columnas):
            ws.cell(row=f, column=columna_inicio + j, value=row[col])


def escribir_lista_plana(
    ruta_plantilla,
    ruta_salida,
    hoja: str,
    fila_inicio: int,
    tabla: pd.DataFrame,
    columnas: list[str],
    columna_inicio: int = COL_A_COD_DEPARTAMENTO,
) -> None:
    """Wrapper de un solo uso - ver docstring de `escribir_cuadro`."""
    origen = ruta_salida if ruta_salida.exists() else ruta_plantilla
    wb = openpyxl.load_workbook(origen)
    _escribir_lista_plana_en_wb(wb, hoja, fila_inicio, tabla, columnas, columna_inicio)
    ruta_salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta_salida)
